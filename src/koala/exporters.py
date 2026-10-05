"""All formats share the same citation rendering and mandatory disclosure."""
import html
import re
from pathlib import Path
from . import DISCLAIMER, is_book, disclaimer_for
from .core import audit, cited_ids, validate_sources
from .citations import MARKER, STYLE_FILES, render_csl, rich_html, spans

FORMATS = ('docx', 'html', 'txt', 'rtf', 'pdf')


def bibliography(s):
    authors = '; '.join(s['authors'])
    venue = s.get('container') or s.get('publisher') or ''
    detail = ', '.join(str(s[k]) for k in ('volume', 'issue', 'pages') if s.get(k))
    doi = f" https://doi.org/{s['doi']}" if s.get('doi') else ''
    publication = ', '.join(str(x) for x in (venue, detail) if x)
    return f"{authors} ({s['year']}). {s['title']}." + (f' {publication}.' if publication else '') + doi


def blocks(article, abstract_only=False):
    sources = validate_sources(article.get('sources', []))
    by_id = {s['id']: s for s in sources}
    body = '\n'.join(s['text'] for s in article.get('sections', []))
    order = list(dict.fromkeys(cited_ids(body)))
    if set(order) - set(by_id) or '[@' in MARKER.sub('', body):
        raise ValueError('Export refused: unknown or malformed citation IDs. Run koala audit.')
    numbers = {key: i + 1 for i, key in enumerate(order)}
    style = article.get('brief', {}).get('citation_style', 'mla')
    csl_references = []
    rendered = None
    if style in STYLE_FILES and not abstract_only:
        paragraphs = [p.strip() for section in article.get('sections', []) for p in section['text'].split('\n\n') if p.strip()]
        values, csl_references = render_csl(paragraphs, sources, style)
        rendered = iter(values)
    # Disambiguate same author/year entries deterministically.
    labels, groups, suffixes = {}, {}, {}
    for key in order:
        s = by_id[key]
        name = s['authors'][0] if s['authors'] else s['title']
        if len(s['authors']) > 1:
            name += ' et al.'
        label = f"{name}, {s['year']}"
        groups.setdefault(label, []).append(key)
    for label, keys in groups.items():
        for i, key in enumerate(keys):
            suffix = ''
            if len(keys) > 1:
                n = i + 1
                while n:
                    n, remainder = divmod(n - 1, 26)
                    suffix = chr(97 + remainder) + suffix
            suffixes[key] = suffix
            labels[key] = (f'[{numbers[key]}]' if style == 'numeric' else f'({label}{suffix})')
    def render(text):
        return next(rendered) if rendered is not None else MARKER.sub(lambda m: labels[m[1]], text)
    result = [('title', article['title']), ('disclosure', disclaimer_for(article)), ('heading', 'Abstract'),
              ('paragraph', article['abstract'])]
    if abstract_only:
        return result
    book = is_book(article)
    chapters = article.get('plan', {}).get('chapters', [])
    if book:
        from .books import chapter_label
        result.append(('heading', 'Contents'))
        result.extend(('contents', chapter_label(c, chapters)) for c in chapters)
        if order:
            result.append(('contents', 'References'))
    reference_heading = 'Works Cited' if style == 'mla' else 'References'
    if book and order:
        result[-1] = ('contents', reference_heading)
    previous_chapter = None
    for section in article.get('sections', []):
        if book:
            index = section.get('chapter_index')
            if type(index) is not int or not 0 <= index < len(chapters):
                raise ValueError('Book section has an invalid chapter assignment.')
            if index != previous_chapter:
                result.append(('chapter', chapter_label(chapters[index], chapters)))
                previous_chapter = index
        result.append(('subheading' if book else 'heading', section['heading']))
        result.extend(('paragraph', render(p.strip())) for p in section['text'].split('\n\n') if p.strip())
    if order:
        result.append(('chapter' if book else 'heading', reference_heading))
    if rendered is not None:
        result.extend(('reference', text) for text in csl_references)
        return result
    reference_order = order if style == 'numeric' else sorted(order, key=lambda k: (
        '; '.join(by_id[k]['authors']).casefold(), str(by_id[k]['year']), by_id[k]['title'].casefold()))
    for key in reference_order:
        prefix = f'[{numbers[key]}] ' if style == 'numeric' else ''
        reference = dict(by_id[key], year=str(by_id[key]['year']) + (suffixes[key] if style != 'numeric' else ''))
        result.append(('reference', prefix + bibliography(reference)))
    return result


def rtf_escape(text):
    out = []
    for char in text:
        n = ord(char)
        if char in '\\{}':
            out.append('\\' + char)
        elif char == '\n':
            out.append('\\line ')
        elif n < 128:
            out.append(char)
        else:
            # RTF uses signed UTF-16 code units, including surrogate pairs.
            encoded = char.encode('utf-16-le')
            for i in range(0, len(encoded), 2):
                unit = int.from_bytes(encoded[i:i+2], 'little')
                out.append(f'\\u{unit if unit < 32768 else unit - 65536}?')
    return ''.join(out)


def write_format(path, content, fmt, pdf_font=None):
    if fmt == 'txt':
        path.write_text('\n\n'.join(t for _, t in content) + '\n', encoding='utf-8')
    elif fmt == 'html':
        tags = {'title': 'h1', 'heading': 'h2', 'chapter': 'h2', 'subheading': 'h3'}
        body = '\n'.join(f'<{tags.get(k,"p")} class="{k}">{rich_html(t)}</{tags.get(k,"p")}>' for k, t in content)
        path.write_text('<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{html.escape(content[0][1])}</title><style>'
            'body{max-width:760px;margin:4rem auto;padding:0 1.5rem;font:18px/1.7 Georgia,serif;color:#171717}'
            'h1{font-size:2rem;line-height:1.2}h2{font-size:1.3rem;margin-top:2rem}'
            '.disclosure{font-size:.9rem;color:#444}.reference{padding-left:1.5rem;text-indent:-1.5rem;overflow-wrap:anywhere}'
            '@media print{body{margin:0;max-width:none;font-size:12pt}h2,h3{break-after:avoid}.chapter{break-before:page}}'
            '</style></head><body><article>' + body + '</article></body></html>', encoding='utf-8')
    elif fmt == 'rtf':
        lines = ['{\\rtf1\\ansi\\deff0\\uc1{\\fonttbl{\\f0 Times New Roman;}}\\margl1440\\margr1440\\fs24']
        for kind, text in content:
            if kind == 'chapter':
                lines.append('\\page')
            control = '\\b\\fs32 ' if kind in ('title', 'chapter') else '\\b\\fs28 ' if kind in ('heading', 'subheading') else '\\fs24 '
            lines.append('{\\pard\\sa180 ' + ('\\li432\\fi-432 ' if kind == 'reference' else '') + control + ''.join(('{\\i '+rtf_escape(t)+'}') if italic else rtf_escape(t) for t,italic in spans(text)) + '\\par}')
        path.write_text('\n'.join(lines) + '\n}', encoding='ascii')
    elif fmt == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml.ns import qn
        doc = Document()
        # Some bundled templates contain themed fonts and a title border.
        for border in doc.styles.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        for style in doc.styles:
            for fonts in style.element.xpath('.//w:rFonts'):
                for name in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme'):
                    fonts.attrib.pop(qn('w:' + name), None)
        for sec in doc.sections:
            sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
        normal = doc.styles['Normal']
        normal.font.name = 'Times New Roman'
        normal.font.size = Pt(12)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(8)
        for name in ('Title', 'Heading 1', 'Heading 2'):
            doc.styles[name].font.name = 'Times New Roman'
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for kind, text in content:
            p = doc.add_paragraph('', 'Title' if kind == 'title' else 'Heading 1' if kind in ('heading', 'chapter') else 'Heading 2' if kind == 'subheading' else 'Normal')
            for fragment, italic in spans(text):
                p.add_run(fragment).italic = italic
            if kind == 'chapter':
                p.paragraph_format.page_break_before = True
            if kind == 'reference':
                p.paragraph_format.left_indent = Inches(.3)
                p.paragraph_format.first_line_indent = Inches(-.3)
                p.paragraph_format.line_spacing = 1.15
            if kind == 'disclosure':
                for run in p.runs:
                    run.italic = True
        if any(kind == 'chapter' for kind, _ in content):
            from docx.oxml import OxmlElement
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            for section in doc.sections:
                footer = section.footer.paragraphs[0]
                footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                field = OxmlElement('w:fldSimple')
                field.set(qn('w:instr'), 'PAGE')
                footer._p.append(field)
        doc.core_properties.title = content[0][1]
        doc.core_properties.comments = next(t for k, t in content if k == 'disclosure')
        doc.save(str(path))
    elif fmt == 'pdf':
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        chars = ''.join(text for _, text in content)
        required = {ord(c) for c in chars if ord(c) > 32}
        import hashlib
        def font_name(path):
            return 'KOALAFont' + hashlib.sha256(str(Path(path).resolve()).encode()).hexdigest()[:16]
        font = 'Times-Roman'
        if pdf_font:
            font = font_name(pdf_font)
            pdfmetrics.registerFont(TTFont(font, str(pdf_font)))
        else:
            candidates = [Path('/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'),
                          Path('/Library/Fonts/Georgia.ttf'),
                          Path('/System/Library/Fonts/Supplemental/Georgia.ttf'),
                          Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf'),
                          Path('/Library/Fonts/Arial Unicode.ttf'),
                          Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
            for candidate in candidates:
                if candidate.exists():
                    try:
                        candidate_font = TTFont(font_name(candidate), str(candidate))
                    except Exception:
                        continue
                    if not required.issubset(candidate_font.face.charToGlyph):
                        continue
                    pdfmetrics.registerFont(candidate_font)
                    font = font_name(candidate)
                    break
        if font != 'Times-Roman':
            regular = Path(pdf_font) if pdf_font else candidate
            italic_path = regular.with_name(regular.stem+' Italic.ttf')
            if regular.name == 'DejaVuSerif.ttf': italic_path = regular.with_name('DejaVuSerif-Italic.ttf')
            if italic_path.exists():
                italic_font = TTFont(font+'Italic', str(italic_path))
                italic_name = font
                if required.issubset(italic_font.face.charToGlyph):
                    pdfmetrics.registerFont(italic_font)
                    italic_name = font+'Italic'
                pdfmetrics.registerFontFamily(font, normal=font, italic=italic_name, bold=font, boldItalic=italic_name)
            else:
                pdfmetrics.registerFontFamily(font, normal=font, italic=font, bold=font, boldItalic=font)
        # Fail clearly for missing glyphs, instead of silently printing black squares.
        chars = ''.join(text for _, text in content)
        registered = pdfmetrics.getFont(font)
        if hasattr(registered.face, 'charToGlyph'):
            missing = {c for c in chars if ord(c) > 32 and ord(c) not in registered.face.charToGlyph}
        else:
            missing = set()
            for c in chars:
                try:
                    c.encode('cp1252')
                except UnicodeEncodeError:
                    missing.add(c)
        if missing:
            raise ValueError('PDF font lacks required characters: ' + ', '.join(
                f'{c!r} (U+{ord(c):04X})' for c in sorted(missing)) +
                '. Your manuscript is saved. Export DOCX/HTML, or supply --pdf-font with a suitable TrueType font.')
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = font
            style.textColor = colors.black
        styles['Normal'].fontSize = 11
        styles['Normal'].leading = 16
        styles['Normal'].spaceAfter = 9
        from reportlab.lib.styles import ParagraphStyle
        styles.add(ParagraphStyle('Reference',parent=styles['Normal'],leftIndent=22,firstLineIndent=-22))
        story = []
        for kind, text in content:
            if kind == 'chapter':
                story.append(PageBreak())
            style = styles['Title' if kind == 'title' else 'Heading1' if kind == 'chapter' else 'Heading2' if kind == 'heading' else 'Heading3' if kind == 'subheading' else 'Reference' if kind == 'reference' else 'Normal']
            story.append(Paragraph(rich_html(text).replace('\n', '<br/>'), style))
            if kind == 'title':
                story.append(Spacer(1, 10))
        def page_number(canvas, doc):
            canvas.setFont(font, 9)
            canvas.drawRightString(A4[0]-60, 35, str(doc.page))
        SimpleDocTemplate(str(path), pagesize=A4, rightMargin=60, leftMargin=60,
            topMargin=60, bottomMargin=55, title=content[0][1], author='KOALA').build(
                story, onFirstPage=page_number, onLaterPages=page_number)
    else:
        raise ValueError(f'Unsupported output format: {fmt}')


def export(article, folder, formats=FORMATS, abstract_only=False, pdf_font=None, strict=False):
    if any(fmt not in FORMATS for fmt in formats):
        raise ValueError(f'Formats must be drawn from {", ".join(FORMATS)}.')
    issues=audit(article)['issues'] if strict else []
    issues=[i for i in issues if not i.startswith(('Citation shortfall:', 'Required authors'))]
    if strict and issues:
        raise ValueError('Strict export refused: resolve audit issues first.')
    content = blocks(article, abstract_only)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    stem = 'abstract' if abstract_only else 'manuscript' if is_book(article) else 'article'
    for fmt in dict.fromkeys(formats):
        dest = folder / f'{stem}.{fmt}'
        tmp = dest.with_name(dest.stem + '.tmp.' + fmt)
        try:
            write_format(tmp, content, fmt, pdf_font)
            tmp.replace(dest)
        finally:
            tmp.unlink(missing_ok=True)
        paths.append(dest)
    return paths
